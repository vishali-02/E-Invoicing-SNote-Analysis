"""
write_results.py <excel_path>

Reads note_data.json and writes all scraped data back into the Excel workbook
produced by init_excel.py.

Columns updated (by header name):
  C: Component
  D: Description
  E: Link                (as a clickable hyperlink)
  F: SP109 Supported     (Yes = green, No = red conditional formatting)
  G: Manual Activity (S4CORE 109)
  H: Activity Timing (Pre/Post)
  I: Attachment
"""
import sys
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

STATE_FILE = os.path.join(os.path.dirname(__file__), "note_data.json")

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL   = PatternFill("solid", fgColor="FFC7CE")
GREEN_FONT = Font(color="276221")
RED_FONT   = Font(color="9C0006")
LINK_FONT  = Font(color="0563C1", underline="single")

# Header names must match HEADERS list in init_excel.py exactly
OUTPUT_HEADERS = [
    "Component",
    "Description",
    "Link",
    "SP109 Supported",
    "Manual Activity (S4CORE 109)",
    "Activity Timing (Pre/Post)",
    "Attachment",
]


def load_state():
    if not os.path.exists(STATE_FILE):
        print("ERROR: note_data.json not found. Run the analysis first.")
        sys.exit(1)
    with open(STATE_FILE, encoding="utf-8") as f:
        return json.load(f)


def build_col_map(ws):
    """Return {header_name: column_index} for all known headers."""
    col_map = {}
    for col in range(1, (ws.max_column or 1) + 1):
        val = ws.cell(row=1, column=col).value
        if val:
            col_map[str(val).strip()] = col
    return col_map


def main():
    if len(sys.argv) < 2:
        print("Usage: python write_results.py <excel_path>")
        sys.exit(1)

    excel_path = sys.argv[1]
    state = load_state()
    results = state.get("results", {})

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    col_map = build_col_map(ws)

    note_col = col_map.get("Note Number", 2)

    # Build row-index map from note numbers already in the sheet
    note_row = {}
    for row_idx in range(2, ws.max_row + 1):
        val = ws.cell(row=row_idx, column=note_col).value
        if val:
            note_row[str(val).strip()] = row_idx

    updated = 0
    skipped = 0

    for note, data in results.items():
        if note not in note_row:
            skipped += 1
            continue
        row_idx = note_row[note]

        def _set(header, value):
            c = col_map.get(header)
            if c:
                ws.cell(row=row_idx, column=c, value=value)

        _set("Component",   data.get("component", ""))
        _set("Description", data.get("description", ""))

        # Link — write as hyperlink
        link_col = col_map.get("Link")
        if link_col:
            link = data.get("link", f"https://me.sap.com/notes/{note}")
            cell = ws.cell(row=row_idx, column=link_col, value=link)
            cell.hyperlink = link
            cell.font = LINK_FONT

        _set("SP109 Supported",             data.get("sp109", ""))
        _set("Manual Activity (S4CORE 109)", data.get("manual", ""))
        _set("Activity Timing (Pre/Post)",   data.get("timing", ""))
        _set("Attachment",                   data.get("attachment", ""))
        updated += 1

    # Conditional formatting on SP109 Supported column
    sp109_col = col_map.get("SP109 Supported")
    if sp109_col and ws.max_row > 1:
        col_letter = ws.cell(row=1, column=sp109_col).column_letter
        range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"Yes"'], fill=GREEN_FILL, font=GREEN_FONT),
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"No"'], fill=RED_FILL, font=RED_FONT),
        )

    wb.save(excel_path)
    print(f"Done. Updated {updated} rows, skipped {skipped}.")
    print(f"Saved: {excel_path}")


if __name__ == "__main__":
    main()
