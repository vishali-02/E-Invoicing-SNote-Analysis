"""
init_excel.py <txt_path> <excel_path>

Reads a text file of SAP note numbers (one per line, or comma/space separated),
creates (or overwrites) an Excel workbook with all required columns, and populates
column A with a sequential row number and column B with the note numbers in order.

Columns created:
  A: S.No
  B: Note Number
  C: Component
  D: Description
  E: Link
  F: SP109 Supported
  G: Manual Activity (S4CORE 109)
  H: Activity Timing (Pre/Post)
  I: Attachment

All headers are styled with a blue header bar (white bold text).
"""

import sys
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS = [
    "S.No",
    "Note Number",
    "Component",
    "Description",
    "Link",
    "SP109 Supported",
    "Manual Activity (S4CORE 109)",
    "Activity Timing (Pre/Post)",
    "Attachment",
]

COLUMN_WIDTHS = {
    "S.No": 7,
    "Note Number": 15,
    "Component": 22,
    "Description": 50,
    "Link": 40,
    "SP109 Supported": 18,
    "Manual Activity (S4CORE 109)": 28,
    "Activity Timing (Pre/Post)": 24,
    "Attachment": 20,
}

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def parse_notes(txt_path):
    """Extract all 7-digit SAP note numbers from the text file, preserving order and deduping."""
    with open(txt_path, encoding="utf-8", errors="replace") as f:
        content = f.read()
    seen = set()
    notes = []
    for match in re.finditer(r"\b(\d{7})\b", content):
        n = match.group(1)
        if n not in seen:
            seen.add(n)
            notes.append(n)
    return notes


def main():
    if len(sys.argv) < 3:
        print("Usage: python init_excel.py <txt_path> <excel_path>")
        sys.exit(1)

    txt_path = sys.argv[1]
    excel_path = sys.argv[2]

    notes = parse_notes(txt_path)
    if not notes:
        print("WARNING: No 7-digit note numbers found in the text file.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OSS Notes"

    # Write headers
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        col_letter = cell.column_letter
        ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(header, 18)

    ws.row_dimensions[1].height = 30

    # Write note numbers
    for i, note in enumerate(notes, start=1):
        ws.cell(row=i + 1, column=1, value=i)           # S.No
        ws.cell(row=i + 1, column=2, value=note)         # Note Number

    wb.save(excel_path)
    print(f"Created Excel with {len(notes)} notes: {excel_path}")
    print(f"Notes: {notes}")

    import json
    print("JSON:" + json.dumps({"notes": notes, "excel_path": excel_path}))


if __name__ == "__main__":
    main()
