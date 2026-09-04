"""
setup_columns.py <excel_path>

Ensures the four output columns exist in row 1 of the active sheet.
Idempotent — safe to run multiple times.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

REQUIRED_HEADERS = [
    "SP109 Supported",
    "Manual Activity",
    "Prerequisites",
    "Root Note",
]

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

COLUMN_WIDTHS = {
    "SP109 Supported": 18,
    "Manual Activity": 18,
    "Prerequisites": 40,
    "Root Note": 15,
}


def find_or_add_header(ws, header_name):
    """Return the column index (1-based) of the header, creating it if absent."""
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        if cell.value and str(cell.value).strip() == header_name:
            return col
    # Not found — append after last column
    new_col = max_col + 1
    cell = ws.cell(row=1, column=new_col, value=header_name)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGN
    col_letter = cell.column_letter
    ws.column_dimensions[col_letter].width = COLUMN_WIDTHS.get(header_name, 20)
    return new_col


def main():
    if len(sys.argv) < 2:
        print("Usage: python setup_columns.py <excel_path>")
        sys.exit(1)

    path = sys.argv[1]
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    col_map = {}
    for header in REQUIRED_HEADERS:
        col_idx = find_or_add_header(ws, header)
        col_map[header] = col_idx
        print(f"  '{header}' -> column {col_idx}")

    wb.save(path)
    print(f"Saved: {path}")
    print("Column map:", col_map)


if __name__ == "__main__":
    main()
