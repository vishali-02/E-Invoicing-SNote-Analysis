"""
write_results.py <excel_path>

Reads note_state.json and writes enriched data back into the Excel file.

- Original rows (already in the sheet): updated in-place.
- Prerequisite-only rows (discovered during walk, not in original sheet):
  appended at the bottom with a light-yellow background to distinguish them.

Columns written (must already exist from setup_columns.py):
  SP109 Supported | Manual Activity | Prerequisites | Root Note
"""
import sys
import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

STATE_FILE = os.path.join(os.path.dirname(__file__), "note_state.json")

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL   = PatternFill("solid", fgColor="FFC7CE")
GREEN_FONT = Font(color="276221")
RED_FONT   = Font(color="9C0006")
PREREQ_ROW_FILL = PatternFill("solid", fgColor="FFFF99")  # light yellow for appended rows

REQUIRED_HEADERS = ["SP109 Supported", "Manual Activity", "Prerequisites", "Root Note"]


def load_state():
    if not os.path.exists(STATE_FILE):
        print("ERROR: note_state.json not found. Run the note analysis first.")
        sys.exit(1)
    with open(STATE_FILE, encoding="utf-8") as f:
        return json.load(f)


def get_col_map(ws):
    col_map = {}
    for col in range(1, (ws.max_column or 1) + 1):
        val = ws.cell(row=1, column=col).value
        if val and str(val).strip() in REQUIRED_HEADERS:
            col_map[str(val).strip()] = col
    return col_map


def get_note_col(ws):
    """Return the column index of the Note Number column (usually column B = 2)."""
    for col in range(1, (ws.max_column or 1) + 1):
        val = ws.cell(row=1, column=col).value
        if val and str(val).strip().lower() in ("note number", "note no", "note"):
            return col
    return 2  # fallback to column B


def write_result(ws, row_idx, col_map, data):
    sp109_col   = col_map.get("SP109 Supported")
    manual_col  = col_map.get("Manual Activity")
    prereq_col  = col_map.get("Prerequisites")
    root_col    = col_map.get("Root Note")

    if sp109_col:
        ws.cell(row=row_idx, column=sp109_col, value=data.get("sp109", ""))
    if manual_col:
        ws.cell(row=row_idx, column=manual_col, value=data.get("manual", ""))
    if prereq_col:
        prereqs = data.get("prereqs", [])
        ws.cell(row=row_idx, column=prereq_col, value=", ".join(prereqs))
    if root_col:
        ws.cell(row=row_idx, column=root_col, value=data.get("root", ""))


def main():
    if len(sys.argv) < 2:
        print("Usage: python write_results.py <excel_path>")
        sys.exit(1)

    path = sys.argv[1]
    state = load_state()
    results = state.get("results", {})

    wb = openpyxl.load_workbook(path)
    ws = wb.active

    col_map  = get_col_map(ws)
    note_col = get_note_col(ws)

    if not col_map:
        print("ERROR: Output columns not found. Run setup_columns.py first.")
        sys.exit(1)

    # Build map of note -> row index for existing rows
    existing_rows = {}
    for row_idx in range(2, ws.max_row + 1):
        note_val = ws.cell(row=row_idx, column=note_col).value
        if note_val:
            existing_rows[str(note_val).strip()] = row_idx

    updated = 0
    appended = 0

    for note, data in results.items():
        if note in existing_rows:
            row_idx = existing_rows[note]
            write_result(ws, row_idx, col_map, data)
            updated += 1
        else:
            # Append as a new row (prerequisite discovered during walk)
            new_row = ws.max_row + 1
            # Write note number in the note column
            ws.cell(row=new_row, column=note_col, value=note)
            # Highlight the entire row in light yellow
            for col in range(1, (ws.max_column or 1) + 1):
                ws.cell(row=new_row, column=col).fill = PREREQ_ROW_FILL
            write_result(ws, new_row, col_map, data)
            # Add SAP for Me link if column 4 exists (based on original sheet pattern)
            link_col = 4
            if ws.cell(row=1, column=link_col).value and "link" in str(ws.cell(row=1, column=link_col).value).lower():
                link = f"https://me.sap.com/notes/{note}"
                cell = ws.cell(row=new_row, column=link_col, value=link)
                cell.hyperlink = link
                cell.font = Font(color="0563C1", underline="single")
            appended += 1

    # Conditional formatting on SP109 column
    sp109_col = col_map.get("SP109 Supported")
    if sp109_col:
        col_letter = ws.cell(row=1, column=sp109_col).column_letter
        last_row = ws.max_row
        range_ref = f"{col_letter}2:{col_letter}{last_row}"
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"Yes"'], fill=GREEN_FILL, font=GREEN_FONT)
        )
        ws.conditional_formatting.add(
            range_ref,
            CellIsRule(operator="equal", formula=['"No"'], fill=RED_FILL, font=RED_FONT)
        )

    wb.save(path)
    print(f"Done. Updated {updated} existing rows, appended {appended} prerequisite rows.")
    print(f"Saved: {path}")


if __name__ == "__main__":
    main()
